import csv

import pytest
from openpyxl import load_workbook

from pdf_tables_to_xlsx.cli import main
from pdf_tables_to_xlsx.normalize import clean_number, normalize_table

T1 = [["Name", "Qty", "Price"], ["Apple", "3", "1.50"], ["Pear", "4", "2.00"]]
T2 = [["City", "Pop"], ["Hanoi", "8"], ["Hue", "1"]]


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_xlsx_basic(tmp_path, make_pdf, capsys):
    pdf = make_pdf(tmp_path / "sample.pdf", [[T1]])
    out = tmp_path / "out.xlsx"
    assert main([str(pdf), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["sample_p1_t1"]
    assert sheet_rows(wb.active) == T1
    assert "page 1: 1 table" in capsys.readouterr().out


def test_clean_numbers(tmp_path, make_pdf):
    table = [["Item", "Amount", "Pct", "Price"], ["A", "1,234.50", "45%", "$9.99"], ["B", "(120)", "10%", "$1.00"]]
    pdf = make_pdf(tmp_path / "n.pdf", [[table]])
    out = tmp_path / "n.xlsx"
    assert main([str(pdf), "-o", str(out), "--clean-numbers"]) == 0
    rows = sheet_rows(load_workbook(out).active)
    assert rows[0] == table[0]
    assert rows[1][1:] == pytest.approx([1234.5, 0.45, 9.99])
    assert rows[2][1] == -120
    assert all(isinstance(v, (int, float)) for v in rows[1][1:])


def test_csv_mode(tmp_path, make_pdf):
    pdf = make_pdf(tmp_path / "sample.pdf", [[T1]])
    out = tmp_path / "outdir"
    assert main([str(pdf), "--format", "csv", "-o", str(out)]) == 0
    files = sorted(p.name for p in out.iterdir())
    assert files == ["sample_p1_t1.csv"]
    with (out / files[0]).open(newline="", encoding="utf-8-sig") as f:
        assert list(csv.reader(f)) == T1


def test_page_range(tmp_path, make_pdf):
    pdf = make_pdf(tmp_path / "sample.pdf", [[T1], [T2]])
    out = tmp_path / "p2.xlsx"
    assert main([str(pdf), "-o", str(out), "--pages", "2"]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["sample_p2_t1"]
    assert sheet_rows(wb.active) == T2


def test_scanned_no_text(tmp_path, make_pdf, capsys):
    pdf = make_pdf(tmp_path / "blank.pdf", [[]])
    out = tmp_path / "x.xlsx"
    assert main([str(pdf), "-o", str(out)]) != 0
    err = capsys.readouterr().err
    assert "no text layer" in err and "scanned" in err
    assert not out.exists()


def test_bad_inputs(tmp_path, capsys):
    out = tmp_path / "x.xlsx"
    assert main([str(tmp_path / "missing.pdf"), "-o", str(out)]) != 0
    assert "not found" in capsys.readouterr().err
    fake = tmp_path / "fake.pdf"
    fake.write_text("hello")
    assert main([str(fake), "-o", str(out)]) != 0
    assert "Not a PDF" in capsys.readouterr().err
    assert not out.exists()


def test_overwrite_protection(tmp_path, make_pdf, capsys):
    pdf = make_pdf(tmp_path / "sample.pdf", [[T1]])
    out = tmp_path / "out.xlsx"
    out.write_bytes(b"keep")
    assert main([str(pdf), "-o", str(out)]) != 0
    assert out.read_bytes() == b"keep"
    assert "--force" in capsys.readouterr().err
    assert main([str(pdf), "-o", str(out), "--force"]) == 0
    assert load_workbook(out).sheetnames == ["sample_p1_t1"]


def test_header_cleanup_in_pdf(tmp_path, make_pdf):
    table = [["Col", "Col", ""], ["a", "b", "c"]]
    pdf = make_pdf(tmp_path / "h.pdf", [[table]])
    out = tmp_path / "h.xlsx"
    assert main([str(pdf), "-o", str(out)]) == 0
    assert sheet_rows(load_workbook(out).active)[0] == ["Col", "Col_2", "Column_3"]


def test_normalize_table_unit():
    raw = [[" Col\n1 ", None, "Col", "", "Col"], [None] * 5, ["x", None, "y", "", "z"]]
    assert normalize_table(raw) == [["Col 1", "Col", "Col_2"], ["x", "y", "z"]]
    assert normalize_table(raw, header=False)[0] == ["Col 1", "Col", "Col"]
    assert normalize_table([[None, ""]]) == []


@pytest.mark.parametrize("text,expected", [
    ("1,234.50", 1234.5), ("(120)", -120), ("45%", 0.45), ("$9.99", 9.99),
    ("-$5", -5), ("1.234,50", 1234.5), ("EUR 10", 10), ("42", 42),
])
def test_clean_number(text, expected):
    assert clean_number(text) == pytest.approx(expected)


@pytest.mark.parametrize("text", ["abc", "007", "12,5", "2024-01-05", "$", "", "1,23"])
def test_clean_number_left_as_text(text):
    assert clean_number(text) == text
