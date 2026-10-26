"""Output writers (xlsx / csv) and unique naming."""
import csv
import re
from pathlib import Path


def unique_name(source: str, page: int, index: int, used: set) -> str:
    """Build a sheet/file name like file_p1_t2 (<=31 chars, unique)."""
    stem = re.sub(r"[\[\]:*?/\\]", "_", source) or "table"
    suffix = f"_p{page}_t{index}"
    base = stem[: 31 - len(suffix)] + suffix
    name, n = base, 2
    while name.lower() in used:
        tag = f"_{n}"
        name, n = base[: 31 - len(tag)] + tag, n + 1
    used.add(name.lower())
    return name


def write_xlsx(path: Path, named: list) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in named:
        ws = wb.create_sheet(name)
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                if value == "":
                    continue
                cell = ws.cell(r, c, value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"  # never create formulas from PDF text
    wb.save(path)


def write_csv(path: Path, rows: list) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)
